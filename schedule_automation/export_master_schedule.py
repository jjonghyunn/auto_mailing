"""
export_master_schedule.py
2026-10-13  Jonghyun Park w/ Claude  — Auto MASTER 'api수기제외' 칼럼이 O 인 site 를 매 회차 읽어 자동 제외 (EXCLUDE_SOURCE_XLSX),
                                       최신 메일 회차 CSV 는 매 회차 다시 계산 (제외 목록 변경이 바로 반영되도록)
2026-10-13  Jonghyun Park w/ Claude  — 검수용 5번째 칼럼 source_subs 추가 (site 가 매칭된 고객 파일 C열 Subs 원본값, 매칭 실패면 빈칸)
                                       + 수기 제외 site 상수 EXCLUDE_SITES (한 줄에 site_code 하나)
(이전 버전 이력은 git history / GitHub Releases 참조 — 헤더에는 최근 2개 항목만 남긴다)

고객 법인 일정 파일(`1.고객 법인 일정 파일/yymmdd_hhmm_*.xlsx`)을 직접 읽어 site_code 별 Start/End 를 계산하고
`campaign, site_code, start_date, end_date, source_subs` 5칼럼 CSV 로 저장한다.
- source_subs = 그 site 의 일정을 가져온 고객 파일 C열(Subs) 원본값 — 검수용. 매칭 실패(일정 없음)면 빈칸.

- 파일명 = schedule_<yymmdd>_<hhmm>.csv — 고객 파일명 앞 수신일시 prefix(check_mail_attachment.py 가 붙임) 그대로
  → 법인 일정 파일의 메일 시각과 1:1 매칭. 아직 CSV 가 없는 고객 파일 + **최신 고객 파일**을 처리한다
  (최신 회차는 매번 다시 계산해 내용이 바뀌었을 때만 덮어씀. 지난 회차는 그 시점 스냅샷으로 남김 — 다시 뽑으려면 --all).
- 법인 → site 매핑 = 같은 폴더 `site_subs_map.csv` (site_code, subs, region, country — 행 순서 = CSV 출력 순서).
  Auto 파일 `태깅기획site_code` 시트에서 뽑은 것. 매핑이 바뀌면 이 CSV 를 고치거나 `--build-map` 으로 다시 뽑는다.
- site 별 계산은 Auto 파일 `고객법인일정파일` 시트 O~S 열 수식 + MASTER H/I 수식을 그대로 옮겼다 (아래 '수식 이식' 참조).
- 미참여는 start 에 'X', 일정 없음(매칭 실패)은 빈칸 — MASTER 와 같은 값.
- 제외 site = 상단 EXCLUDE_SITES (수기) ∪ EXCLUDE_SOURCE_XLSX 의 MASTER 'api수기제외' 칼럼이 O 인 site (자동, 매 회차 읽음).
  CSV 와 --check 대조 양쪽에서 뺀다. 제외 소스 파일을 못 읽으면 그 회차는 통째로 SKIP (제외가 빠진 CSV 를 쓰지 않도록).

옵션:
  --all        : CSV 가 이미 있는 고객 파일도 다시 계산 (내용이 다르면 덮어씀)
  --check      : 최신 고객 파일 결과를 Auto MASTER 와 site 별로 대조 (Auto D1 스탬프가 같은 파일일 때만)
  --build-map  : Auto 파일 `태깅기획site_code` 시트에서 site_subs_map.csv 를 다시 만든다
"""

import csv
import io
import re
import sys
import datetime as dt
from pathlib import Path
import openpyxl

# ════ 사용자가 바꿔야 하는 부분 ════

# ─── 캠페인 식별자 (CSV campaign 칼럼값) ───
CAMPAIGN_ID = "campaign_id"

# ─── 수기 제외 site (한 줄에 site_code 하나, 대소문자 무시. 빈 줄·# 주석 무시) ───
EXCLUDE_SITES = """
"""

# ─── 자동 제외 소스 (이 파일 MASTER 시트의 'api수기제외' 칼럼이 O 인 site 를 매 회차 제외) ───
# None 이면 자동 제외 끔 (EXCLUDE_SITES 만 사용)
EXCLUDE_SOURCE_XLSX = Path(
    r"C:\Users\user_name\OneDrive - company_name"
    r"\Project_team_name - 1 company_name - 02 part_name"
    r"\part_name\2026\# CAMPAIGN_PROJECTS\02. CAMPAIGN NAME\02. SCHEDULE"
    r"\26_Schedule(Auto).xlsx"
)
EXCLUDE_SOURCE_SHEET      = "MASTER"
EXCLUDE_SOURCE_HEADER_ROW = 4
EXCLUDE_SOURCE_SITE_COL   = "G"            # Site Code
EXCLUDE_SOURCE_FLAG_COL   = "J"            # api수기제외
EXCLUDE_SOURCE_FLAG_HDR   = "api수기제외"   # 헤더가 이 값이 아니면 구조가 바뀐 것 → 그 회차 SKIP
EXCLUDE_SOURCE_FLAG_MARK  = "O"            # 대소문자 무시

# ─── 경로 ───
BASE = Path(
    r"C:\Users\user_name\OneDrive - company_name"
    r"\Project_team_name - 1 company_name - 02 part_name"
    r"\part_name\2026\# CAMPAIGN_PROJECTS\02. CAMPAIGN NAME\02. SCHEDULE"
)
SOURCE_FOLDER   = BASE / "1.고객 법인 일정 파일"
OUTPUT_DIR      = BASE / "3.main_raw_schedule"
OUTPUT_PREFIX   = "schedule_"
OUTPUT_ENCODING = "utf-8-sig"          # Excel 에서 바로 열리게 BOM 포함
SITE_MAP_CSV    = Path(__file__).resolve().parent / "site_subs_map.csv"

# ─── 고객 파일 읽기 (update_schedule.py 와 같은 규칙 — Auto 시트에 붙는 모양을 그대로 재현) ───
SRC_FIRST_ROW    = 3      # 헤더 포함 읽기 시작 행
SRC_FIRST_COL    = 2      # B열
SRC_LAST_COL     = 11     # K열
NOTE_HEADER      = "Note" # 이 헤더 칼럼은 J열 위치로 옮긴다 (update_schedule 과 동일)
NOTE_TARGET_COL  = 10     # J열
NOTE_HEADER_ROWS = 5
FILL_MERGED_COLS = [5]    # E열 Participation 세로 병합 채움
SUBS_HEADER      = "Subs" # C열이 이 값인 행 = 헤더행. 데이터는 그 다음 행부터 (Auto 수식 범위 C4:C200 과 같음)

# ─── Auto 파일 (--check 대조 / --build-map 에만 사용) ───
AUTO_FILE_NAME    = "26_Schedule(Auto).xlsx"
MASTER_SHEET      = "MASTER"
MASTER_HEADER_ROW = 4
MASTER_COLS       = {"site": "G", "start": "H", "end": "I"}
STAMP_SHEET       = "고객법인일정파일"
STAMP_CELL        = "D1"
TAG_SHEET         = "태깅기획site_code"
TAG_HEADER_ROW    = 3
TAG_COLS          = {"no": "B", "region": "C", "subs": "D", "country": "E", "site": "F"}

# ─── 특수 표기 (Auto 수식에 박혀 있던 값) ───
HQ_SUBS_KR    = "한총"    # 고객 파일 C열 법인명이 이 값(앞부분)이면 본사(HQ_SUBS_CLEAN) 법인으로 본다
HQ_SUBS_CLEAN = "hq"
PART_YES      = "O"
PART_NO       = "X"

# ─── 출력 칼럼 ───
OUTPUT_COLUMNS = ["campaign", "site_code", "start_date", "end_date", "source_subs"]
MAP_COLUMNS    = ["site_code", "subs", "region", "country"]
DATE_FORMAT    = "%Y-%m-%d"

# ════ 내부 사용 ════

MAIL_PREFIX_PAT = re.compile(r'^(\d{6})_(\d{4})_')   # check_mail_attachment.py 저장명 규칙
FLAG_ALL, FLAG_CHECK, FLAG_BUILD_MAP = "--all", "--check", "--build-map"


def _col_idx(letter: str) -> int:
    return openpyxl.utils.column_index_from_string(letter)


def _fmt(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime(DATE_FORMAT)
    if v is None:
        return ""
    return str(v).strip()


def _ci(v) -> str:
    """Excel COUNTIF/MATCH 처럼 대소문자 무시 비교용 키."""
    return "" if v is None else str(v).casefold()


def _blank(v) -> bool:
    return v is None or (isinstance(v, str) and v == "")


# ── 고객 파일 → Auto `고객법인일정파일` 시트에 붙는 행 (B~K) ──────────────
def read_customer_rows(src: Path) -> list:
    """update_schedule.py 의 읽기 규칙(빈 행 스킵 / E열 병합 채움 / Note → J열)을 그대로 따라 B~K 행 리스트를 만든다."""
    wb = openpyxl.load_workbook(src, data_only=True)
    try:
        ws = wb.worksheets[0]
        fill = {}
        for mr in ws.merged_cells.ranges:
            if mr.min_col != mr.max_col or mr.min_col not in FILL_MERGED_COLS or mr.min_row == mr.max_row:
                continue
            top = ws.cell(row=mr.min_row, column=mr.min_col).value
            if top is None:
                continue
            for r in range(mr.min_row + 1, mr.max_row + 1):
                fill[(r, mr.min_col)] = top
        rows = []
        for row in ws.iter_rows(min_row=SRC_FIRST_ROW, min_col=SRC_FIRST_COL, max_col=SRC_LAST_COL):
            if all(c.value is None for c in row):
                continue
            vals = []
            for c in row:
                v = c.value if c.value is not None else fill.get((c.row, c.column))
                vals.append(v.date() if isinstance(v, dt.datetime) else v)
            rows.append(vals)
    finally:
        wb.close()

    if NOTE_HEADER:
        note_idx = next((i for r in rows[:NOTE_HEADER_ROWS] for i, v in enumerate(r)
                         if isinstance(v, str) and v.strip().lower() == NOTE_HEADER.lower()), None)
        tgt = NOTE_TARGET_COL - SRC_FIRST_COL
        if note_idx is not None and note_idx != tgt:
            for r in rows:
                r.insert(tgt, r.pop(note_idx))

    # 헤더행(C열 'Subs') 다음부터가 데이터 — Auto 수식이 보는 C4:C200 과 같은 범위
    hdr = next((i for i, r in enumerate(rows) if isinstance(r[1], str) and r[1].strip() == SUBS_HEADER), None)
    return rows[hdr + 1:] if hdr is not None else rows


# ── 수식 이식 ────────────────────────────────────────────────
# B=0 C=1 D=2 E=3 F=4 G=5 H=6 (고객 행 인덱스)
IB, IC, IE, IF_, IH = 0, 1, 3, 4, 6


def ar_clean(row) -> str:
    """O열 AR_clean: Region/Total/빈 행은 '', 그 외 C열의 '-' 또는 ' ' 앞부분. HQ_SUBS_KR 이면 HQ_SUBS_CLEAN."""
    c = row[IC]
    if not _blank(row[IB]) or _blank(c) or c == "Total":
        return ""
    c = str(c)
    if "-" in c:
        v = c[:c.find("-")]
    elif " " in c:
        v = c[:c.find(" ")]
    else:
        v = c
    return HQ_SUBS_CLEAN if v == HQ_SUBS_KR else v


def ar_key(subs: str, site: str, c_keys: set, o_to_c: dict, c_first: dict):
    """P열 AR_key: 태깅 site 에 대응하는 고객 C열 값 (3단계 우선순위 + fallback). 못 찾으면 ''."""
    if not site:
        return ""
    base = site.split("_", 1)[0]
    k1 = f"{subs}-{base.upper()}"
    if _ci(k1) in c_keys:
        return k1
    if "_" in site:
        lang = site.split("_", 1)[1][:20]
        k2 = f"{subs} ({base.title()}_{lang.upper()})"
        if _ci(k2) in c_keys:
            return k2
    if _ci(subs) in c_keys:
        return subs
    if "_" in site:
        return ""
    if _ci(subs) in o_to_c:                      # O열(정제값) 매칭 → 그 행의 C 값
        return o_to_c[_ci(subs)]
    if "-" in subs:
        head = subs[:subs.find("-")]
        return c_first.get(_ci(head), "")
    return ""                                    # (수식의 SCA 분기는 '_' 있는 site 에서만 의미가 있어 위에서 이미 '' — 도달 불가)


def compute_sites(cust_rows: list, site_map: list) -> list:
    """[(site, start, end, source_subs)] — start/end 는 MASTER H/I 와 같은 값, source_subs 는 매칭된 고객 C열 원본."""
    c_first, o_to_c = {}, {}
    for r in cust_rows:
        if not _blank(r[IC]):
            c_first.setdefault(_ci(r[IC]), r)
        o = ar_clean(r)
        if o:
            o_to_c.setdefault(_ci(o), r[IC])
    c_keys = set(c_first)
    c_row = {k: r for k, r in c_first.items()}

    out = []
    for m in site_map:
        key = ar_key(m["subs"], m["site_code"], c_keys, o_to_c,
                     {k: r[IC] for k, r in c_first.items()})
        start = end = src_subs = ""
        if str(key).strip():
            r = c_row.get(_ci(key))
            src_subs = r[IC] if r is not None else ""
            v = str(r[IE]).upper() if r is not None and r[IE] is not None else ""
            part = PART_YES if "O" in v else (PART_NO if "X" in v else "")
            if part == PART_NO:
                start = PART_NO
            elif part == PART_YES and r is not None:
                start, end = r[IF_], r[IH]
        out.append((m["site_code"], start, end, src_subs))
    return out


# ── site 매핑 CSV ───────────────────────────────────────────
def manual_excluded_sites() -> set:
    return {ln.split("#", 1)[0].strip().casefold() for ln in EXCLUDE_SITES.splitlines()} - {""}


def source_excluded_sites() -> set:
    """EXCLUDE_SOURCE_XLSX MASTER 'api수기제외' = O 인 site. 읽기 실패·헤더 불일치는 예외로 올린다 (호출부가 회차 SKIP)."""
    if EXCLUDE_SOURCE_XLSX is None:
        return set()
    wb = openpyxl.load_workbook(EXCLUDE_SOURCE_XLSX, data_only=True, read_only=True)
    try:
        ws = wb[EXCLUDE_SOURCE_SHEET]
        si, fi = _col_idx(EXCLUDE_SOURCE_SITE_COL) - 1, _col_idx(EXCLUDE_SOURCE_FLAG_COL) - 1
        rows = list(ws.iter_rows(min_row=EXCLUDE_SOURCE_HEADER_ROW, max_col=max(si, fi) + 1, values_only=True))
    finally:
        wb.close()
    hdr = _fmt(rows[0][fi]) if rows else ""
    if hdr != EXCLUDE_SOURCE_FLAG_HDR:
        raise ValueError(f"{EXCLUDE_SOURCE_SHEET}!{EXCLUDE_SOURCE_FLAG_COL}{EXCLUDE_SOURCE_HEADER_ROW} 헤더가 "
                         f"{EXCLUDE_SOURCE_FLAG_HDR!r} 아님 ({hdr!r}) — 칼럼 위치 확인")
    mark = EXCLUDE_SOURCE_FLAG_MARK.casefold()
    return {_fmt(r[si]).casefold() for r in rows[1:]
            if not _blank(r[si]) and _fmt(r[fi]).casefold() == mark}


def load_site_map() -> list:
    if not SITE_MAP_CSV.exists():
        raise FileNotFoundError(f"site 매핑 CSV 없음: {SITE_MAP_CSV} — `python {Path(__file__).name} {FLAG_BUILD_MAP}` 로 만든다")
    with SITE_MAP_CSV.open(encoding=OUTPUT_ENCODING, newline="") as f:
        rows = [r for r in csv.DictReader(f) if (r.get("site_code") or "").strip()]
    if not rows:
        raise ValueError(f"site 매핑 CSV 가 비어 있음: {SITE_MAP_CSV}")
    for r in rows:
        r["site_code"], r["subs"] = r["site_code"].strip(), (r.get("subs") or "").strip()
    return rows


def build_site_map() -> int:
    auto = BASE / AUTO_FILE_NAME
    wb = openpyxl.load_workbook(auto, data_only=True, read_only=True)
    try:
        ws = wb[TAG_SHEET]
        idx = {k: _col_idx(v) - 1 for k, v in TAG_COLS.items()}
        rows = list(ws.iter_rows(min_row=TAG_HEADER_ROW + 1, max_col=max(idx.values()) + 1, values_only=True))
    finally:
        wb.close()
    out = [{"site_code": str(r[idx["site"]]).strip(), "subs": _fmt(r[idx["subs"]]),
            "region": _fmt(r[idx["region"]]), "country": _fmt(r[idx["country"]])}
           for r in rows if not _blank(r[idx["no"]]) and not _blank(r[idx["site"]])]
    with SITE_MAP_CSV.open("w", encoding=OUTPUT_ENCODING, newline="") as f:
        w = csv.DictWriter(f, fieldnames=MAP_COLUMNS, lineterminator="\n")
        w.writeheader()
        w.writerows(out)
    print(f"[매핑 생성] {SITE_MAP_CSV.name} ({len(out)} site) ← {auto.name} '{TAG_SHEET}'")
    return 0


# ── 출력 ─────────────────────────────────────────────────────
def csv_text(results) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(OUTPUT_COLUMNS)
    for site, start, end, src_subs in results:
        w.writerow([CAMPAIGN_ID, site, _fmt(start), _fmt(end), _fmt(src_subs)])
    return buf.getvalue()


def write_csv(out_file: Path, text: str, n: int) -> None:
    if out_file.exists():
        if out_file.read_text(encoding=OUTPUT_ENCODING).replace("\r\n", "\n") == text:
            print(f"[SKIP] 이미 저장됨: {out_file.name}")
            return
        action = "갱신"
    else:
        action = "저장"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_file.write_text(text, encoding=OUTPUT_ENCODING, newline="")
    print(f"[{action}] {out_file.name} ({n} site)")


def check_against_master(src: Path, results, excl: set) -> None:
    """Auto MASTER(답안지)와 site 별 대조 — Auto D1 스탬프가 이 고객 파일일 때만."""
    auto = BASE / AUTO_FILE_NAME
    try:
        wb = openpyxl.load_workbook(auto, data_only=True, read_only=True)
    except (OSError, FileNotFoundError) as e:
        print(f"[대조 생략] Auto 파일을 읽을 수 없음 ({type(e).__name__}: {e})")
        return
    try:
        stamp = wb[STAMP_SHEET][STAMP_CELL].value
        idx = {k: _col_idx(v) - 1 for k, v in MASTER_COLS.items()}
        rows = list(wb[MASTER_SHEET].iter_rows(min_row=MASTER_HEADER_ROW + 1,
                                               max_col=max(idx.values()) + 1, values_only=True))
    finally:
        wb.close()
    if str(stamp or "").strip() != src.name:
        print(f"[대조 생략] Auto 에 반영된 소스가 다름 (Auto {stamp!r} / 이번 {src.name!r})")
        return
    master = {str(r[idx["site"]]).strip(): (_fmt(r[idx["start"]]), _fmt(r[idx["end"]]))
              for r in rows if not _blank(r[idx["site"]]) and str(r[idx["site"]]).strip().casefold() not in excl}
    mine = {s: (_fmt(a), _fmt(b)) for s, a, b, _ in results}
    diff = [(s, mine.get(s), master.get(s)) for s in sorted(set(mine) | set(master)) if mine.get(s) != master.get(s)]
    if diff:
        print(f"[대조 불일치] {len(diff)} site (계산값 / MASTER)")
        for s, a, b in diff:
            print(f"   {s}: {a} / {b}")
    else:
        print(f"[대조 일치] MASTER 와 {len(mine)} site 모두 같음")


def main() -> int:
    args = sys.argv[1:]
    if FLAG_BUILD_MAP in args:
        return build_site_map()

    try:
        site_map = load_site_map()
    except (OSError, ValueError) as e:
        print(f"[ERROR] {e}")
        return 1

    manual = manual_excluded_sites()
    try:
        auto_excl = source_excluded_sites()
    except (OSError, KeyError, ValueError) as e:
        # 잠김 / 시트·칼럼 구조 변경 — 제외가 빠진 CSV 를 쓰지 않도록 이번 회차는 건너뛴다
        print(f"[SKIP] 제외 소스를 읽을 수 없음 ({type(e).__name__}: {e})")
        return 0
    excl = manual | auto_excl
    if excl:
        print(f"[제외] 수기 {sorted(manual)} / 소스 api수기제외 {sorted(auto_excl)}")
    unknown = excl - {m["site_code"].casefold() for m in site_map}
    if unknown:
        print(f"[경고] 제외 site 중 매핑표에 없는 site: {sorted(unknown)}")
    site_map = [m for m in site_map if m["site_code"].casefold() not in excl]

    sources = sorted((f for f in SOURCE_FOLDER.glob("*.xlsx")
                      if not f.name.startswith("~$") and MAIL_PREFIX_PAT.match(f.name)), key=lambda f: f.name)
    if not sources:
        print(f"[SKIP] 수신일시 prefix 가 붙은 고객 법인 일정 파일이 아직 없습니다: {SOURCE_FOLDER}")
        return 0

    for src in sources:
        m = MAIL_PREFIX_PAT.match(src.name)
        out_file = OUTPUT_DIR / f"{OUTPUT_PREFIX}{m.group(1)}_{m.group(2)}.csv"
        is_latest = src == sources[-1]
        if out_file.exists() and FLAG_ALL not in args and not is_latest:
            continue
        try:
            results = compute_sites(read_customer_rows(src), site_map)
            print(f"[소스] {src.name}")
            write_csv(out_file, csv_text(results), len(results))
        except (OSError, ValueError, IndexError) as e:
            # OneDrive 동기화 중 잠김 / 포맷이 다른 파일 — 다음 회차가 재시도한다
            print(f"[SKIP] {src.name} 처리 실패 ({type(e).__name__}: {e})")
            continue
        if FLAG_CHECK in args and is_latest:
            check_against_master(src, results, excl)
    return 0


if __name__ == "__main__":
    sys.exit(main())
