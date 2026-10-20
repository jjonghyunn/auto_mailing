r"""
sync_urls_to_raw_url_sheet.py
2026-10-20  Jonghyun Park w/ Claude  — 신규: urls.csv 에만 있는 TY URL 을 Auto 파일 `RAW_URL,PageName` 시트에 추가

왜 있나
  `<일정 폴더>\3.main_raw_schedule\urls.csv` = 이번 캠페인(TY) 법인 페이지 URL 목록 (사람이 추가 — added_at = 추가 시각).
  (같은 csv 를 페이지 캡처 쪽 동기화 스크립트도 읽어 캡처 대상에 넣는다 — 이 repo 밖 도구)
  Auto 파일은 `RAW_URL,PageName` 시트의 TY 행을 MASTER O열(This Year URL)·P열(Page Name)이 FILTER 로 끌어가므로,
  csv 에만 있고 시트에 없는 URL 은 MASTER 에 안 잡힌다 → 그 차이를 시트에 자동으로 메운다.
  (MASTER O열은 TY URL 점검 도구의 원천이기도 하다 — 이 repo 밖 도구)

동작
1. urls.csv (campaign_id, sitecode, url, added_at) 를 읽는다. 같은 URL 은 첫 행만
2. Auto 파일 `RAW_URL,PageName` 시트에서 **B열 = TY** 인 행의 D열 URL 목록과 대조
   - 비교는 앞뒤 공백·`http(s)://`·끝 `/`·대소문자를 무시한다 (쓸 때는 urls.csv 값 그대로)
   - LY 행은 보지도 건드리지도 않는다
3. 시트에 없는 URL 만 마지막 행 아래에 B=TY / C=sitecode / D=url 로 추가 (E pagename 은 비워 둔다 — 사람이 채운다)
   - **추가만 한다.** csv 에서 지운 URL 을 시트에서 지우거나, 기존 행을 고치지 않는다
   - Excel COM(별도 인스턴스)으로 열어 셀에 직접 쓰고 전체 재계산 후 저장
     — openpyxl 로 재저장하지 않으니 G~J 의 UNIQUE/FILTER 동적 배열 수식이 그대로 남는다
   - Excel 이 "피호출자가 호출을 거부했습니다"(RPC_E_CALL_REJECTED) 로 거절하면 처음부터 다시 연다 (COM_RETRIES)
   - 저장 후 다시 읽어 전부 들어갔는지 확인한다
4. 추가할 게 없으면 Auto 파일을 열지 않고 끝난다 (OneDrive 동기화 부담 없음)
5. 누가 파일을 열어 두었거나(잠금 파일 `~$…` / 읽기 전용으로 열림) 저장·검증에 실패하면
   추가하려던 목록만 urls.csv 옆 `PENDING_XLSX` 로 저장한다 — **항상 같은 이름, 최신 1개만** (덮어쓰기).
   반영에 성공하거나 추가할 게 없으면 헤더만 남긴 빈 파일로 둔다 (파일은 항상 있다 — 행이 있으면 = 아직 반영 안 된 목록).

⚠ PENDING_XLSX 는 Auto 파일이 외부 참조로 읽는다 (2026-10-20 사용자가 연결)
  MASTER O열 = RAW_URL TY 행에 그 site URL 이 없으면 `[urls_pending_raw_url.xlsx]pending!C:C` 를 B열(site) 로 FILTER.
  → 시트명 `pending`, B=site_code, C=url 위치를 바꾸지 말 것. 파일을 지우지 말 것 (링크가 깨진다).

스케줄
  작업 스케줄러 `campaign_raw_url_sync` — 20분마다 매시 :14·:34·:54 (C:\Python314\pythonw.exe).
  같은 Auto 파일을 쓰는 `campaign_schedule_update`(:04·:24·:44), 읽는 `campaign_export_master_schedule`(:10…)
  및 그 밖에 Auto 를 읽는 작업과 안 겹치도록 비어 있던 분을 골랐다. 시각을 옮길 땐 그 작업들을 피할 것.
  등록 명령: 같은 폴더 `create_schtasks_v2.txt` 맨 위.

한계
  - 웹 Excel / 공동 편집으로 열려 있으면 `~$` 잠금 파일이 안 생겨 사용 중을 감지 못 할 수 있다
    → 저장은 되지만 OneDrive 가 충돌 사본을 만들 수 있다. 저장 후 검증에서 누락되면 pending 으로 남는다.
  - sitecode 값은 검증하지 않는다 (csv 값 그대로 C열에 들어간다).

실행:  python sync_urls_to_raw_url_sheet.py            # 반영
       python sync_urls_to_raw_url_sheet.py --dry-run  # 추가 대상만 출력 (파일 무변경)
       필요 패키지: openpyxl, pywin32 (Excel 설치 PC)
로그:  같은 폴더 `LOG_FILE` (실행마다 덮어쓰기 — pythonw 는 stdout 을 버리므로)
종료코드: 0 = 반영 완료 / 추가 없음 / Auto 파일 읽기 실패로 SKIP(다음 회차 재시도), 2 = 미반영(pending 저장), 1 = 예외
"""

import csv
import re
import sys
import time
import datetime as dt
from pathlib import Path
import openpyxl
import win32com.client

# ════ 사용자가 바꿔야 하는 부분 ════
# ─── 경로 ───
BASE = Path(
    r"C:\Users\user_name\OneDrive - company_name"
    r"\Project_team_name - 1 company_name - 02 part_name"
    r"\part_name\2026\# CAMPAIGN_PROJECTS\02. CAMPAIGN NAME\02. SCHEDULE"
)
AUTO_FILE_NAME = "26_Schedule(Auto).xlsx"
URLS_CSV       = BASE / "3.main_raw_schedule" / "urls.csv"
PENDING_XLSX   = URLS_CSV.parent / "urls_pending_raw_url.xlsx"   # 반영 실패 시 추가 대상 목록 (최신 1개만)
PENDING_SHEET  = "pending"   # ⚠ Auto MASTER O열이 `pending!B:C` 를 외부 참조 — 시트명·열 순서(B=site, C=url) 고정

# ─── 대상 시트 ───
TARGET_SHEET = "RAW_URL,PageName"
HEADER_ROWS  = 1          # 1행 = 헤더
COL_NO       = 1          # A = 행 번호 수식
COL_YEAR     = 2          # B = This/Last Year
COL_SITE     = 3          # C = site_code
COL_URL      = 4          # D = url
LAST_DATA_COL = 6         # B~F 중 값이 있는 마지막 행 = 데이터 끝 (G~J 는 수식 spill 이라 제외)
TY_LABEL     = "TY"
NO_FORMULA   = '=IF(C{r}="","",ROW()-1)'   # A열이 비어 있는 새 행에만 채운다

# ─── Excel COM 재시도 ───
COM_RETRIES        = 5    # "피호출자가 호출을 거부했습니다"(Excel 이 아직 바쁨) 일 때 처음부터 다시 연다
COM_RETRY_WAIT_SEC = 5

# ─── urls.csv 칼럼 ───
CSV_SITE_COL  = "sitecode"
CSV_URL_COL   = "url"
CSV_ADDED_COL = "added_at"

# ════ 내부 사용 ════
AUTO_FILE = BASE / AUTO_FILE_NAME
LOCK_FILE = BASE / f"~${AUTO_FILE_NAME}"   # Excel 이 열면 생기는 소유자 파일
LOG_FILE  = Path(__file__).with_name("sync_urls_to_raw_url_sheet_last.log")
DRY_RUN   = "--dry-run" in sys.argv
_SCHEME   = re.compile(r"^https?://", re.I)
RPC_E_CALL_REJECTED = -2147418111

_log_lines = []


def log(msg=""):
    print(msg)
    _log_lines.append(msg)


def flush_log():
    try:
        LOG_FILE.write_text("\n".join(_log_lines) + "\n", encoding="utf-8")
    except OSError:
        pass


def norm(url):
    return _SCHEME.sub("", str(url).strip()).rstrip("/").lower()


def read_csv_urls():
    """urls.csv → [(sitecode, url, added_at)] (같은 URL 은 첫 행만)"""
    out, seen = [], set()
    with open(URLS_CSV, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            url = (row.get(CSV_URL_COL) or "").strip()
            if not url or norm(url) in seen:
                continue
            seen.add(norm(url))
            out.append(((row.get(CSV_SITE_COL) or "").strip(), url, (row.get(CSV_ADDED_COL) or "").strip()))
    return out


def read_sheet():
    """(TY URL 정규화 set, 마지막 데이터 행)"""
    wb = openpyxl.load_workbook(AUTO_FILE, read_only=True, data_only=True)
    try:
        ws = wb[TARGET_SHEET]
        ty, last = set(), HEADER_ROWS
        for r_idx, r in enumerate(ws.iter_rows(min_row=HEADER_ROWS + 1, max_col=LAST_DATA_COL, values_only=True),
                                  start=HEADER_ROWS + 1):
            if any(v not in (None, "") for v in r[COL_YEAR - 1:LAST_DATA_COL]):
                last = r_idx
            if str(r[COL_YEAR - 1] or "").strip().upper() == TY_LABEL and r[COL_URL - 1]:
                ty.add(norm(r[COL_URL - 1]))
        return ty, last
    finally:
        wb.close()


def write_pending(rows, reason):
    """반영 못 한 목록을 PENDING_XLSX 로 덮어쓰기 (rows 가 비면 헤더만 있는 빈 파일)
    열: A=This/Last Year · B=site_code · C=url · D=urls.csv added_at · E=reason · F=created_at
    (B·C 는 Auto MASTER O열 외부 참조가 읽는다 — 순서 바꾸지 말 것)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = PENDING_SHEET
    ws.append(["This/Last Year", "site_code", "url", "urls.csv added_at", "reason", "created_at"])
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for site, url, added in rows:
        ws.append([TY_LABEL, site, url, added, reason, now])
    for col, width in zip("ABCDEF", (14, 12, 60, 20, 40, 20)):
        ws.column_dimensions[col].width = width
    try:
        wb.save(PENDING_XLSX)
        if rows:
            log(f"[백업] 추가 대상 {len(rows)}건을 {PENDING_XLSX.name} 에 저장했습니다 ({reason})")
        else:
            log(f"[정리] 반영할 목록이 없어 {PENDING_XLSX.name} 를 빈 파일(헤더만)로 바꿨습니다")
    except PermissionError:
        log(f"[경고] {PENDING_XLSX.name} 가 열려 있어 저장하지 못했습니다")


def clear_pending():
    """반영할 게 없으면 PENDING_XLSX 를 헤더만 있는 빈 파일로 둔다 (지우지 않음).
    이미 빈 파일이면 다시 쓰지 않는다 — 20분마다 OneDrive 에 새 버전이 올라가지 않게."""
    if PENDING_XLSX.exists():
        try:
            wb = openpyxl.load_workbook(PENDING_XLSX, read_only=True)
            empty = wb.active.max_row <= 1
            wb.close()
            if empty:
                return
        except Exception:
            pass                                  # 깨진 파일이면 새로 쓴다
    write_pending([], "")


def write_with_excel(rows, last_row):
    """Excel 이 바빠 호출을 거부하면 COM_RETRIES 번까지 처음부터 다시 시도 (저장 전 실패는 Close(False) 로 버려져 중복 없음)"""
    for attempt in range(1, COM_RETRIES + 1):
        reason, rejected = _write_once(rows, last_row)
        if not rejected or attempt == COM_RETRIES:
            return reason
        log(f"[재시도] Excel 이 호출을 거부 — {COM_RETRY_WAIT_SEC}초 뒤 다시 엽니다 ({attempt}/{COM_RETRIES})")
        time.sleep(COM_RETRY_WAIT_SEC)


def _write_once(rows, last_row):
    """Excel COM 으로 행 추가 → 재계산 → 저장. (실패 사유 문자열 또는 None, 호출 거부 여부)"""
    excel = win32com.client.DispatchEx("Excel.Application")   # 별도 인스턴스 — 사용자가 열어 둔 Excel 에 붙지 않게
    excel.Visible = False
    excel.DisplayAlerts = False
    wb_com = None
    try:
        wb_com = excel.Workbooks.Open(str(AUTO_FILE.resolve()), 0, False)
        if wb_com.ReadOnly:
            return "다른 사용자가 사용 중 (읽기 전용으로 열림)", False
        ws = wb_com.Worksheets(TARGET_SHEET)
        for i, (site, url, _) in enumerate(rows, start=1):
            r = last_row + i
            if not ws.Cells(r, COL_NO).Formula:
                ws.Cells(r, COL_NO).Formula = NO_FORMULA.format(r=r)
            ws.Cells(r, COL_YEAR).Value = TY_LABEL
            ws.Cells(r, COL_SITE).Value = site
            ws.Cells(r, COL_URL).Value = url
        excel.CalculateFull()
        wb_com.Save()
        return None, False
    except Exception as e:                       # COM 오류(잠김·OneDrive 동기화 중 등) — 다음 실행이 재시도
        rejected = bool(getattr(e, "args", None)) and e.args[0] == RPC_E_CALL_REJECTED
        return f"Excel 저장 실패 ({type(e).__name__}: {e})", rejected
    finally:
        # 지연 바인딩이면 Close/Quit 가 TypeError·AttributeError 로 죽는다 (update_schedule.py 와 같은 처리)
        if wb_com is not None:
            try:
                wb_com.Close(SaveChanges=False)
            except (TypeError, AttributeError):
                pass
            except Exception:
                pass
        try:
            excel.Quit()
        except Exception:
            pass


def main():
    log(f"[실행] {dt.datetime.now():%Y-%m-%d %H:%M:%S}{'  (dry-run)' if DRY_RUN else ''}")
    csv_rows = read_csv_urls()
    try:
        ty, last_row = read_sheet()
    except (OSError, KeyError) as e:
        log(f"[SKIP] Auto 파일을 읽을 수 없습니다 ({type(e).__name__}: {e}). 다음 실행 시 재시도합니다")
        return 0
    missing = [r for r in csv_rows if norm(r[1]) not in ty]
    log(f"[대조] urls.csv {len(csv_rows)}건 · 시트 TY URL {len(ty)}건 · 마지막 데이터 행 {last_row} → 추가 대상 {len(missing)}건")
    for site, url, added in missing:
        log(f"   + {site:10s} {url}  (added_at {added})")

    if not missing:
        if not DRY_RUN:
            clear_pending()
        return 0
    if DRY_RUN:
        return 0

    if LOCK_FILE.exists():
        reason = f"파일 사용 중 (잠금 파일 {LOCK_FILE.name})"
    else:
        reason = write_with_excel(missing, last_row)
    if reason is None:
        try:
            ty_after, _ = read_sheet()
            left = [r for r in missing if norm(r[1]) not in ty_after]
            reason = f"저장 후 검증 실패 — {len(left)}건 누락" if left else None
        except (OSError, KeyError) as e:
            reason = f"저장 후 재확인 실패 ({type(e).__name__})"
    if reason is None:
        log(f"[완료] {len(missing)}건을 {TARGET_SHEET} {last_row + 1}~{last_row + len(missing)}행에 추가했습니다")
        clear_pending()
        return 0
    log(f"[미반영] {reason}")
    write_pending(missing, reason)
    return 2


if __name__ == "__main__":
    code = 1
    try:
        code = main()
    except Exception as e:
        log(f"[오류] {type(e).__name__}: {e}")
        raise
    finally:
        flush_log()
    sys.exit(code)
